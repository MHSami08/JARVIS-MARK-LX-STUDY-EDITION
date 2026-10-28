"""
plugins/study_mentor.py — Lightweight Backend Edition.

All primary tasks (Quiz, Flashcards, Mindmap, Problem Solver) are handled by
dedicated action modules in actions/. This module is the backend for notepad
evaluation, session export, and mistake review.

Phase 6: `export` now compiles the whole study session — the notepad plus every
solution / quiz / flashcard set / mindmap shown this session — into a formatted
Microsoft Word document under workspace/exports/. A Markdown copy is written as
a fallback when python-docx is unavailable.
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

try:
    from memory import study_cache
except Exception:
    study_cache = None

try:
    from core import model_router
except Exception:
    model_router = None

_BASE = Path(__file__).resolve().parent.parent
EXPORT_WORKSPACE_DIR = _BASE / "workspace" / "exports"
_SESSION_HISTORY: list[dict] = []      # fallback store (check_notes reviews)


def _lang(text: str) -> str:
    t = str(text or "").lower()
    if any(k in t for k in ["in english", "english version", "ইংরেজি", "ইংলিশে"]):
        return "en"
    return "bn"


def _run_check_notes(player) -> str:
    notes = (player.get_notepad_text() or "").strip() if hasattr(player, "get_notepad_text") else ""
    if not notes:
        return "[STUDY_BOARD] The Study Notepad is empty. Ask the student to press Ctrl+N and write something first."

    lang = _lang(notes)
    player.study_status("Reviewing your notes…" if lang == "en" else "আপনার নোট পরীক্ষা করা হচ্ছে…")
    sys_p = (
        "You are a supportive academic mentor reviewing a student's rough notes. "
        "Format output in Markdown with `# Note Review` and these exact sections:\n"
        "## ✓ What You Got Right\n"
        "## ⚠ Mistakes or Missing Points\n"
        "## ★ Corrected & Polished Version\n"
        "Write in the same language as the notes. Wrap math in LaTeX."
    )
    raw = model_router.call(f"STUDENT'S NOTES:\n{notes[:3500]}", task="study_fast", system=sys_p, timeout=45.0)
    if not raw:
        return "[STUDY_BOARD] Could not evaluate notes right now."

    md = re.sub(r"^```(?:markdown|md)?\s*\n", "", raw.strip(), flags=re.I)
    md = re.sub(r"\n```\s*$", "", md.strip())

    _SESSION_HISTORY.append({"title": "Notepad Review", "markdown": md})
    player.study_solution("Notepad Review", md, "Note Check", reveal="all")
    return f"[STUDY_BOARD] Notepad review is on the board. Give 1-2 sentences of spoken feedback.\nSUMMARY:\n{md[:1500]}"


# ── Phase 6: session aggregation ─────────────────────────────────────────────

def _collect_session(player) -> tuple[str, list[dict]]:
    """Return (notepad_text, [ {title, kind, subject, text}, ... ]).

    Prefers the UI's live board history (Phase 6). Falls back to the module's
    own review log so an older build still exports something useful.
    """
    notes = ""
    try:
        if hasattr(player, "get_notepad_text"):
            notes = (player.get_notepad_text() or "").strip()
    except Exception:
        notes = ""

    items: list[dict] = []
    try:
        if hasattr(player, "get_study_history"):
            for h in (player.get_study_history() or []):
                txt = (h.get("text") or "").strip()
                if txt:
                    items.append({
                        "title": h.get("title") or "Study Item",
                        "kind":  h.get("kind") or "note",
                        "subject": h.get("subject") or "",
                        "text": txt,
                    })
    except Exception:
        items = []

    if not items:                       # legacy fallback
        for h in _SESSION_HISTORY:
            md = (h.get("markdown") or "").strip()
            if md:
                items.append({"title": h.get("title") or "Study Item",
                              "kind": "note", "subject": "", "text": md})
    return notes, items


def _build_markdown(notes: str, items: list[dict]) -> str:
    parts = [f"# JARVIS Study Notes ({datetime.now().strftime('%Y-%m-%d %H:%M')})\n"]
    if notes:
        parts.append(f"## My Notepad\n{notes}\n")
    for it in items:
        head = f"## {it['title']}"
        if it.get("subject"):
            head += f"  _({it['subject']})_"
        parts.append(f"{head}\n{it['text']}\n")
    return "\n\n---\n\n".join(parts)


# ── Phase 6: Markdown → .docx (lightweight, dependency-safe) ─────────────────

_CODE_FONT = "Consolas"


def _add_inline(paragraph, text: str) -> None:
    """Add text to a paragraph, honouring **bold** and `code` spans."""
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            paragraph.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            run = paragraph.add_run(tok[1:-1])
            run.font.name = _CODE_FONT
        else:
            paragraph.add_run(tok)


def _write_markdown_docx(doc, md: str) -> None:
    """Render a small, well-defined subset of Markdown into Word elements."""
    in_code = False
    code_buf: list[str] = []

    def _flush_code():
        if code_buf:
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_buf))
            run.font.name = _CODE_FONT
            code_buf.clear()

    for raw in md.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            if in_code:
                _flush_code()
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_buf.append(raw)
            continue
        if not line.strip():
            doc.add_paragraph()
            continue
        stripped = line.strip()
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
            continue
        if stripped.startswith(("- ", "* ", "• ")):
            indent = max(0, (len(line) - len(line.lstrip())) // 2)
            p = doc.add_paragraph(style="List Bullet")
            if indent:
                p.paragraph_format.left_indent = None
            _add_inline(p, stripped[2:].strip())
            continue
        m = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            _add_inline(p, m.group(1).strip())
            continue
        m = re.match(r"^([★✓⚠\w\s]{1,40}):$", stripped)   # "Label:" pseudo-heading
        if m and len(stripped) < 45:
            p = doc.add_paragraph()
            p.add_run(stripped).bold = True
            continue
        p = doc.add_paragraph()
        _add_inline(p, stripped)
    if in_code:
        _flush_code()


def _export_docx(notes: str, items: list[dict], out_path: Path) -> bool:
    """Write the session to a .docx. Returns True on success."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except Exception:
        return False
    try:
        doc = Document()
        try:
            doc.core_properties.title = "JARVIS Study Notes"
            doc.core_properties.author = "JARVIS"
        except Exception:
            pass
        title = doc.add_heading("JARVIS Study Notes", level=0)
        stamp = doc.add_paragraph()
        run = stamp.add_run(datetime.now().strftime("Exported %Y-%m-%d %H:%M"))
        run.italic = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        if notes:
            doc.add_heading("My Notepad", level=1)
            _write_markdown_docx(doc, notes)

        for it in items:
            doc.add_heading(it["title"] + (f"  ({it['subject']})" if it.get("subject") else ""), level=1)
            _write_markdown_docx(doc, it["text"])

        out_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out_path))
        return True
    except Exception:
        return False


def _run_export(player) -> str:
    notes, items = _collect_session(player)
    if not notes and not items:
        return "[STUDY_BOARD] Nothing to export yet — no notepad text and no board items this session."

    EXPORT_WORKSPACE_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    docx_path = EXPORT_WORKSPACE_DIR / f"JARVIS_Study_Notes_{stamp}.docx"

    if _export_docx(notes, items, docx_path):
        return (f"[STUDY_BOARD] Exported the session to a Word document: {docx_path} "
                f"({len(items)} board item(s) + notepad). Tell the user it is saved in "
                f"workspace/exports.")

    # Fallback: python-docx unavailable or save failed → keep a Markdown copy.
    md_path = EXPORT_WORKSPACE_DIR / f"JARVIS_Study_Notes_{stamp}.md"
    try:
        md_path.write_text(_build_markdown(notes, items), encoding="utf-8")
        return (f"[STUDY_BOARD] Word export unavailable; saved a Markdown copy instead: {md_path}")
    except Exception as e:
        return f"[STUDY_BOARD] Export failed: {e}"


def run(parameters, player=None, **kwargs) -> str:
    action = str((parameters or {}).get("action") or "").lower().strip()

    if action == "review_mistakes":
        mistakes = study_cache.get_mistakes("all", limit=10) if study_cache else []
        if not mistakes:
            return "[STUDY_BOARD] Your Mistake Bank is clean — no saved wrong answers found! Congratulate the student."
        player.study_quiz("Mistake Bank Revision", mistakes, "Revision")
        return f"[STUDY_BOARD] Loaded {len(mistakes)} previously missed questions. Tell the user to press 1-4 to practice them."

    if action == "check_notes":
        return _run_check_notes(player)

    if action == "export":
        return _run_export(player)

    return f"Unknown or deprecated action in study_mentor: {action}"


PLUGIN = {
    "name": "study_mode",
    "description": (
        "Backend engine for the Study Station. 'check_notes' reviews the notepad; "
        "'export' compiles the whole session (notepad + every solution, quiz, "
        "flashcard set and mindmap shown) into a formatted Microsoft Word document "
        "under workspace/exports/; 'review_mistakes' replays the Mistake Bank."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "action": {
                "type": "STRING",
                "description": "check_notes | export | review_mistakes"
            }
        },
        "required": ["action"],
    },
}
